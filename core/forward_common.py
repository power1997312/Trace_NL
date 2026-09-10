from __future__ import annotations
"""
正向追踪矩阵公共核心

供 generate_forward_matrix.py (用户需求→系统需求) 与
generate_forward_sd_matrix.py (系统需求→系统设计) 复用:
- XML级RichText读取(逆向矩阵Excel → 着色runs, 支持多sheet与合并单元格)
- PDF条目提取与引用匹配(精确→子串→章节号前缀→模糊, 4级回退)
- Union着色(锚定基准文本, 字符级取最高优先级 GREEN>BLUE>BLACK)
- Excel输出(表头/行高自适应/富文本单元格/合并单元格)

两入口脚本的差异仅保留在各自的数据构建策略:
- A沿用逆向矩阵已保存的着色runs, 不重新匹配;
- B对每条(上游,下游)重新执行微块匹配, 并补全孤儿上游条目。
"""
import os
import re
import zipfile
import difflib
import xml.etree.ElementTree as ET
from collections import OrderedDict

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.cell.rich_text import CellRichText, TextBlock
from openpyxl.cell.text import InlineFont

from config import (
    HEADER_FILL, COLOR_GREEN, COLOR_BLUE, COLOR_BLACK,
    FONT_NAME, HEADER_FONT_SIZE, CONTENT_FONT_SIZE, ID_FONT_SIZE,
    MatchCategory,
)

NS = 'http://schemas.openxmlformats.org/spreadsheetml/2006/main'

# InlineFont定义 (用于RichText着色)
_IF_GREEN = InlineFont(rFont=FONT_NAME, sz=CONTENT_FONT_SIZE, color=COLOR_GREEN)
_IF_BLUE = InlineFont(rFont=FONT_NAME, sz=CONTENT_FONT_SIZE, color=COLOR_BLUE)
_IF_BLACK = InlineFont(rFont=FONT_NAME, sz=CONTENT_FONT_SIZE, color=COLOR_BLACK)

# 通用样式
_THIN_BORDER = Border(
    left=Side(style='thin'), right=Side(style='thin'),
    top=Side(style='thin'), bottom=Side(style='thin'),
)
_HEADER_FILL = PatternFill(start_color=HEADER_FILL, end_color=HEADER_FILL, fill_type='solid')
_HEADER_FONT = Font(name=FONT_NAME, size=HEADER_FONT_SIZE, bold=True, color='FFFFFF')
_HEADER_ALIGNMENT = Alignment(horizontal='center', vertical='center', wrap_text=True)
_DATA_ALIGNMENT = Alignment(vertical='top', wrap_text=True)


# ============================================================
# 颜色与类别转换
# ============================================================

def hex_to_category(rgb):
    """将RGB颜色字符串转为MatchCategory"""
    if not rgb:
        return MatchCategory.BLACK
    rgb = rgb.upper()
    if len(rgb) == 8:
        rgb = rgb[2:]
    if rgb == '00B050':
        return MatchCategory.GREEN
    elif rgb == '0070C0':
        return MatchCategory.BLUE
    return MatchCategory.BLACK


def category_to_rgb(cat):
    """MatchCategory转RGB字符串"""
    if cat == MatchCategory.GREEN:
        return '00B050'
    elif cat == MatchCategory.BLUE:
        return '0070C0'
    return '000000'


def dominant_category_from_runs(runs):
    """从runs列表判断主色调类别"""
    green_len = 0
    blue_len = 0
    for color, text in runs:
        if not text:
            continue
        cat = hex_to_category(color)
        if cat == MatchCategory.GREEN:
            green_len += len(text)
        elif cat == MatchCategory.BLUE:
            blue_len += len(text)
    if green_len > 0 and green_len >= blue_len:
        return MatchCategory.GREEN
    elif blue_len > 0:
        return MatchCategory.BLUE
    return MatchCategory.BLACK


# ============================================================
# XML级RichText读取
# ============================================================

def read_reverse_matrix_data(xlsx_path):
    """
    从逆向追踪矩阵Excel中读取所有数据(含颜色)

    通过直接解析XML获取富文本颜色信息。
    读取全部sheet(多下游文档时每份一个sheet), 逐sheet处理下游侧合并单元格。

    Returns:
        list[dict]: 每行包含 seq, ds_id, ds_text, ds_runs, ds_category,
                    ds_doc(下游来源文档名, G列), up_doc, up_ref, up_text,
                    up_runs, up_category
    """
    with zipfile.ZipFile(xlsx_path) as z:
        # 读取sharedStrings
        shared = []
        if 'xl/sharedStrings.xml' in z.namelist():
            root = ET.parse(z.open('xl/sharedStrings.xml')).getroot()
            for si in root.findall(f'{{{NS}}}si'):
                runs = si.findall(f'{{{NS}}}r')
                if runs:
                    shared.append([
                        (r.find(f'{{{NS}}}rPr').find(f'{{{NS}}}color').get('rgb', '')
                         if r.find(f'{{{NS}}}rPr') is not None
                            and r.find(f'{{{NS}}}rPr').find(f'{{{NS}}}color') is not None
                         else '',
                         r.find(f'{{{NS}}}t').text
                         if r.find(f'{{{NS}}}t') is not None
                            and r.find(f'{{{NS}}}t').text else '')
                        for r in runs
                    ])
                else:
                    t = si.find(f'{{{NS}}}t')
                    shared.append([('', t.text if t is not None and t.text else '')])

        # 读取所有sheet (逆向追踪矩阵可能有多个sheet, 每份下游文档一个)
        sheet_paths = sorted([n for n in z.namelist()
                              if n.startswith('xl/worksheets/sheet') and n.endswith('.xml')])
        if not sheet_paths:
            raise ValueError("未找到逆向追踪矩阵sheet")

        all_roots = [ET.parse(z.open(sp)).getroot() for sp in sheet_paths]
        rows = []

        for root in all_roots:
            current_ds = None  # 每个sheet重置合并单元格追踪器
            for row_el in root.iter(f'{{{NS}}}row'):
                row_num = int(row_el.get('r', '0'))
                if row_num <= 1:
                    continue

                cells = {}
                for cell_el in row_el.findall(f'{{{NS}}}c'):
                    ref = cell_el.get('r', '')
                    if not ref:
                        continue
                    col_letter = ''.join(c for c in ref if c.isalpha())

                    # 尝试读取inline string
                    is_el = cell_el.find(f'{{{NS}}}is')
                    if is_el is not None:
                        runs_xml = is_el.findall(f'{{{NS}}}r')
                        if runs_xml:
                            runs = []
                            for r in runs_xml:
                                rPr = r.find(f'{{{NS}}}rPr')
                                t = r.find(f'{{{NS}}}t')
                                text = t.text if t is not None and t.text else ''
                                color = ''
                                if rPr is not None:
                                    c = rPr.find(f'{{{NS}}}color')
                                    if c is not None:
                                        color = c.get('rgb', '')
                                runs.append((color, text))
                            full_text = ''.join(text for _, text in runs)
                            category = dominant_category_from_runs(runs)
                            cells[col_letter] = {'text': full_text, 'runs': runs, 'category': category}
                            continue
                        else:
                            t_el = is_el.find(f'{{{NS}}}t')
                            if t_el is not None and t_el.text:
                                cells[col_letter] = {
                                    'text': t_el.text, 'runs': None, 'category': MatchCategory.BLACK,
                                }
                                continue

                    # 尝试读取shared string
                    ct = cell_el.get('t', '')
                    v = cell_el.find(f'{{{NS}}}v')
                    if ct == 's' and v is not None and v.text:
                        sidx = int(v.text)
                        if sidx < len(shared):
                            runs = shared[sidx]
                            full_text = ''.join(text for _, text in runs)
                            category = dominant_category_from_runs(runs)
                            cells[col_letter] = {'text': full_text, 'runs': runs, 'category': category}
                            continue

                    # 普通值
                    if v is not None and v.text:
                        cells[col_letter] = {'text': v.text, 'runs': None, 'category': MatchCategory.BLACK}

                if not cells:
                    continue

                # A=序号, B=下游ID, C=下游内容, D=上游文档, E=上游引用, F=上游内容, G=下游文档(来源)
                seq = cells.get('A', {}).get('text', '')
                ds_id = cells.get('B', {}).get('text', '')
                ds_data = cells.get('C', {})
                up_doc = cells.get('D', {}).get('text', '')
                up_ref = cells.get('E', {}).get('text', '')
                up_data = cells.get('F', {})
                ds_doc = cells.get('G', {}).get('text', '')

                # 更新当前下游条目追踪器 (处理下游侧合并单元格)
                # 仅当ds_id和ds_text同时存在时才更新缓存
                ds_text_val = ds_data.get('text', '')
                if ds_id and ds_text_val:
                    current_ds = {
                        'ds_id': ds_id,
                        'ds_text': ds_text_val,
                        'ds_runs': ds_data.get('runs'),
                        'ds_category': ds_data.get('category', MatchCategory.BLACK),
                    }
                elif ds_id and not ds_text_val and current_ds:
                    # B列有值但C列为空(合并单元格) -> 使用缓存的下游内容
                    ds_data = {
                        'text': current_ds['ds_text'],
                        'runs': current_ds['ds_runs'],
                        'category': current_ds['ds_category'],
                    }
                if not ds_id and current_ds:
                    # B列为空(合并单元格) -> 使用缓存的下游条目号和内容
                    ds_id = current_ds['ds_id']
                    ds_data = {
                        'text': current_ds['ds_text'],
                        'runs': current_ds['ds_runs'],
                        'category': current_ds['ds_category'],
                    }

                if not up_ref:
                    continue

                rows.append({
                    'seq': seq,
                    'ds_id': ds_id,
                    'ds_text': ds_data.get('text', ''),
                    'ds_runs': ds_data.get('runs'),
                    'ds_category': ds_data.get('category', MatchCategory.BLACK),
                    'ds_doc': ds_doc,
                    'up_doc': up_doc,
                    'up_ref': up_ref,
                    'up_text': up_data.get('text', ''),
                    'up_runs': up_data.get('runs'),
                    'up_category': up_data.get('category', MatchCategory.BLACK),
                })

    return rows


# ============================================================
# PDF条目提取与引用匹配
# ============================================================

def normalize_ref(ref):
    """归一化引用字符串: 去空格、统一全半角"""
    ref = ref.strip()
    ref = ref.replace('（', '(').replace('）', ')').replace('：', ':')
    ref = ref.replace('，', ',').replace('；', ';')
    ref = re.sub(r'\s+', '', ref)
    return ref


def extract_pdf_items(pdf_path, doc_name=''):
    """
    从PDF中提取所有需求条目和章节

    - ID型文档: 仅提取需求条目, 不重复提取章节
    - 章节型文档: 提取章节, 标题截断防止正文泄漏

    Returns:
        items: 每项含 key, content, section_number, section_title, item_id
    """
    from core.pdf_parser_adapter import extract_full_text
    from core.requirement_extractor import (
        extract_requirement_items, extract_sections,
        detect_document_type,
    )

    text = extract_full_text(pdf_path)
    doc_type = detect_document_type(text)
    items = []

    if doc_type == "id_based":
        # ID型文档: 仅按需求条目获取, 不再按章节号重复提取
        for item in extract_requirement_items(text):
            items.append({
                'key': item.item_id,
                'content': item.content,
                'section_number': '',
                'section_title': '',
                'item_id': item.item_id,
            })
    else:
        # 章节型文档: 提取章节, 标题截断防止正文泄漏
        _MAX_TITLE_LEN = 12
        for section in extract_sections(text):
            title = section.section_title
            if len(title) > _MAX_TITLE_LEN:
                # 策略1: 在句号/冒号处截断
                period_pos = title.find('。')
                colon_pos = title.find('：')
                cut = -1
                for pos in (period_pos, colon_pos):
                    if 4 <= pos <= _MAX_TITLE_LEN and (cut < 0 or pos < cut):
                        cut = pos
                if cut > 0:
                    title = title[:cut]
                else:
                    # 策略2: 在逗号处截断
                    comma_pos = title.find('，')
                    if 4 <= comma_pos <= _MAX_TITLE_LEN:
                        title = title[:comma_pos]
                    else:
                        title = title[:_MAX_TITLE_LEN]
            items.append({
                'key': f"{section.section_number} {title}",
                'content': section.content,
                'section_number': section.section_number,
                'section_title': title,
                'item_id': '',
            })

    return items


def build_item_key_map(items):
    """
    构建归一化键→条目索引的映射

    支持多种键形式: item_id, section_number, section_number+title(带空格/不带空格), title
    """
    key_map = {}
    for idx, item in enumerate(items):
        keys = []
        if item.get('item_id'):
            keys.append(item['item_id'])
        if item.get('section_number'):
            keys.append(item['section_number'])
            if item.get('section_title'):
                # 同时注册带空格和不带空格两种格式
                keys.append(f"{item['section_number']} {item['section_title']}")
                keys.append(f"{item['section_number']}{item['section_title']}")
        if item.get('section_title'):
            keys.append(item['section_title'])
        keys.append(item['key'])

        for k in keys:
            nk = normalize_ref(k)
            if nk and nk not in key_map:
                key_map[nk] = idx
    return key_map


def match_ref_to_item(ref, key_map, items):
    """
    将引用字符串匹配到PDF条目

    策略: 精确匹配 → 子串匹配 → 章节号前缀匹配 → 模糊匹配
    """
    ref_norm = normalize_ref(ref)

    # 1. 精确匹配
    if ref_norm in key_map:
        return key_map[ref_norm]

    # 2. 子串匹配 (要求双方长度≥5, 且非纯章节号, 避免父节抢匹配)
    if len(ref_norm) >= 5:
        for nk, idx in key_map.items():
            if len(nk) >= 5 and not re.match(r'^\d+(?:\.\d+)*$', nk) \
                    and (ref_norm in nk or nk in ref_norm):
                return idx

    # 3. 章节号前缀匹配
    ref_num_m = re.match(r'^(\d+(?:\.\d+)*)', ref_norm)
    if ref_num_m:
        ref_num = ref_num_m.group(1)
        best_idx = None
        best_len = 0
        for nk, idx in key_map.items():
            key_num_m = re.match(r'^(\d+(?:\.\d+)*)', nk)
            if key_num_m:
                kn = key_num_m.group(1)
                if kn == ref_num:
                    return idx
                if kn.startswith(ref_num) or ref_num.startswith(kn):
                    if len(kn) > best_len:
                        best_len = len(kn)
                        best_idx = idx
        if best_idx is not None:
            return best_idx

    # 4. 模糊匹配
    best_idx = None
    best_ratio = 0.0
    for nk, idx in key_map.items():
        ratio = difflib.SequenceMatcher(None, ref_norm, nk).ratio()
        if ratio > best_ratio:
            best_ratio = ratio
            best_idx = idx
    if best_ratio >= 0.7:
        return best_idx

    return None


# ============================================================
# Union Coloring (组合合并着色)
# ============================================================

def textruns_to_tuples(runs):
    """将 TextRun 列表转为 (rgb_hex, text) 元组列表, 供 union coloring / rich text 复用"""
    out = []
    for r in runs:
        rgb = category_to_rgb(r.category)
        out.append((rgb, r.text))
    return out


def compute_union_coloring(all_runs_list, base_text):
    """
    计算Union着色: 多条追踪关系的同一上游内容按字符取最高优先级颜色

    优先级: GREEN(3) > BLUE(2) > BLACK(1)

    以 base_text 为锚定: 颜色数组长度恒等于 len(base_text),
    runs 超出基准文本长度的部分被忽略, 不足部分保持 BLACK —
    保证输出 runs 的拼接文本与 base_text 严格一致(Excel 渲染不丢字)。
    """
    priority = {MatchCategory.GREEN: 3, MatchCategory.BLUE: 2, MatchCategory.BLACK: 1}
    max_len = len(base_text)
    char_colors = [MatchCategory.BLACK] * max_len

    for runs in all_runs_list:
        if not runs:
            continue
        pos = 0
        for color, text in runs:
            cat = hex_to_category(color)
            for k in range(len(text)):
                if pos + k < max_len and priority.get(cat, 0) > priority.get(char_colors[pos + k], 0):
                    char_colors[pos + k] = cat
            pos += len(text)

    # 转回runs
    runs_out = []
    cur_rgb = None
    cur_text = []
    for i, ch in enumerate(base_text):
        cat = char_colors[i] if i < len(char_colors) else MatchCategory.BLACK
        rgb = category_to_rgb(cat)
        if cur_rgb is None:
            cur_rgb = rgb
            cur_text = [ch]
        elif rgb == cur_rgb:
            cur_text.append(ch)
        else:
            runs_out.append((cur_rgb, ''.join(cur_text)))
            cur_rgb = rgb
            cur_text = [ch]
    if cur_text:
        runs_out.append((cur_rgb, ''.join(cur_text)))
    return runs_out


# ============================================================
# Excel输出
# ============================================================

def make_rich_text_cell(runs, default_text=''):
    """将runs列表转为CellRichText"""
    if not runs:
        return default_text

    blocks = []
    for color, text in runs:
        if not text:
            continue
        cat = hex_to_category(color)
        if cat == MatchCategory.GREEN:
            ifont = _IF_GREEN
        elif cat == MatchCategory.BLUE:
            ifont = _IF_BLUE
        else:
            ifont = _IF_BLACK
        blocks.append(TextBlock(ifont, text))

    if not blocks:
        return default_text
    if len(blocks) == 1:
        return CellRichText(blocks[0])
    return CellRichText(*blocks)


def estimate_visible_lines(text, col_width, font_size=CONTENT_FONT_SIZE):
    """估算文本在指定列宽(Excel字符单位)下 wrap_text 后的可见行数。

    Excel 列宽 1 单位 ≈ 1 个半角字符宽度; 全角/CJK 字符按 2 单位计。
    wrap_text 时文本按列宽自动折行, 显式换行(\\n)强制分段。
    """
    if not text:
        return 1
    chars_per_line = max(4.0, col_width - 2.0)
    total = 0
    for seg in str(text).split('\n'):
        if not seg:
            total += 1
            continue
        w = 0.0
        for ch in seg:
            w += 2.0 if ord(ch) > 0x7F else 1.0
        total += max(1, -(-int(w) // int(chars_per_line)))
    return max(1, total)


def set_row_height(ws, row, texts_with_widths, min_height=15.0,
                   line_height=17.0, pad=4.0):
    """根据单元格内容估算并设置行高, 使 wrap_text 多行可见。"""
    max_lines = 1
    for text, col_width in texts_with_widths:
        if text:
            max_lines = max(max_lines, estimate_visible_lines(text, col_width))
    ws.row_dimensions[row].height = max(min_height, max_lines * line_height + pad)


def set_merged_row_heights(ws, first_row, last_row, texts_with_widths,
                           min_height=15.0, line_height=17.0, pad=4.0):
    """合并单元格区域的行高按行数平均分摊。"""
    max_lines = 1
    for text, col_width in texts_with_widths:
        if text:
            max_lines = max(max_lines, estimate_visible_lines(text, col_width))
    total_height = max(min_height, max_lines * line_height + pad)
    n = max(1, last_row - first_row + 1)
    per_row = max(min_height, total_height / n)
    for r in range(first_row, last_row + 1):
        ws.row_dimensions[r].height = per_row


def set_cell(ws, row, col, value, font=None):
    """设置单元格值和样式"""
    cell = ws.cell(row=row, column=col, value=value)
    if font:
        cell.font = font
    cell.alignment = _DATA_ALIGNMENT
    cell.border = _THIN_BORDER
    return cell


def write_header_row(ws, headers, widths):
    """写入表头行"""
    from openpyxl.utils import get_column_letter
    for col_idx, (header, width) in enumerate(zip(headers, widths), 1):
        cell = ws.cell(row=1, column=col_idx, value=header)
        cell.font = _HEADER_FONT
        cell.fill = _HEADER_FILL
        cell.alignment = _HEADER_ALIGNMENT
        cell.border = _THIN_BORDER
        # 支持超过26列(AA/AB/...)
        ws.column_dimensions[get_column_letter(col_idx)].width = width


def generate_forward_excel(forward_data_by_doc, output_path, doc_val_fn):
    """
    生成正向追踪矩阵Excel(统一输出结构)

    - 每份上游文档一个独立sheet
    - 上游内容(C列)合并, 使用union coloring
    - 下游内容(F列)按追踪关系独立着色
    - 无追踪关系的条目, 下游D/E/F列填"NA"

    Args:
        forward_data_by_doc: {doc_name: OrderedDict{(up_doc, up_ref): entry}}
            entry 含 up_text / up_runs / has_trace / ds_items;
            ds_items 每项含 ds_id / ds_text / ds_runs / ds_doc(可选)
        output_path: 输出Excel路径
        doc_val_fn: fn(ds_item) -> 下游文档名(D列), 由各流水线决定取值来源
    """
    wb = Workbook()
    first_sheet = True

    for doc_name, entries in forward_data_by_doc.items():
        if first_sheet:
            ws = wb.active
            ws.title = doc_name[:31]
            first_sheet = False
        else:
            ws = wb.create_sheet(title=doc_name[:31])

        headers = ['序号', '上游条目号', '上游内容', '下游文档', '下游条目号', '下游内容']
        widths = [6, 22, 55, 20, 22, 55]
        write_header_row(ws, headers, widths)

        row_num = 2
        seq = 1

        for (up_doc, up_ref), data in entries.items():
            ds_items = data.get('ds_items') or []
            has_trace = data.get('has_trace')

            if has_trace and ds_items:
                # 有追踪关系: 一行上游条目对应多行下游条目
                start_row = row_num

                for di in ds_items:
                    # A: 序号
                    set_cell(ws, row_num, 1, seq,
                             Font(name=FONT_NAME, size=ID_FONT_SIZE))

                    # B: 上游条目号
                    set_cell(ws, row_num, 2, up_ref,
                             Font(name=FONT_NAME, size=ID_FONT_SIZE))

                    # C: 上游内容 (union coloring, 合并)
                    up_rich = make_rich_text_cell(data.get('up_runs'), data.get('up_text', ''))
                    ws.cell(row=row_num, column=3, value=up_rich)
                    ws.cell(row=row_num, column=3).alignment = _DATA_ALIGNMENT
                    ws.cell(row=row_num, column=3).border = _THIN_BORDER

                    # D: 下游文档
                    doc_val = doc_val_fn(di)
                    set_cell(ws, row_num, 4, doc_val,
                             Font(name=FONT_NAME, size=ID_FONT_SIZE))

                    # E: 下游条目号
                    set_cell(ws, row_num, 5, di['ds_id'],
                             Font(name=FONT_NAME, size=ID_FONT_SIZE))

                    # F: 下游内容 (按该条追踪关系着色)
                    ds_rich = make_rich_text_cell(di.get('ds_runs'), di.get('ds_text', ''))
                    ws.cell(row=row_num, column=6, value=ds_rich)
                    ws.cell(row=row_num, column=6).alignment = _DATA_ALIGNMENT
                    ws.cell(row=row_num, column=6).border = _THIN_BORDER

                    # 数据行行高(临时按当前行内容估算, 合并组在下方统一分摊覆盖)
                    set_row_height(ws, row_num, [
                        (data.get('up_text', ''), 55),
                        (di.get('ds_text', ''), 55),
                    ])

                    row_num += 1

                end_row = row_num - 1

                # 合并 A, B, C 列 (序号、上游条目号、上游内容)
                if end_row > start_row:
                    for col in [1, 2, 3]:
                        ws.merge_cells(
                            start_row=start_row, start_column=col,
                            end_row=end_row, end_column=col,
                        )
                    # 合并组行高: 以上游内容为基准, 按行数分摊
                    set_merged_row_heights(ws, start_row, end_row, [
                        (data.get('up_text', ''), 55),
                    ])

                seq += 1
            else:
                # 无追踪关系: 下游填NA
                set_cell(ws, row_num, 1, seq,
                         Font(name=FONT_NAME, size=ID_FONT_SIZE))
                set_cell(ws, row_num, 2, up_ref,
                         Font(name=FONT_NAME, size=ID_FONT_SIZE))

                # C: 上游内容 (来自PDF, 无着色)
                set_cell(ws, row_num, 3, data.get('up_text', ''),
                         Font(name=FONT_NAME, size=CONTENT_FONT_SIZE))

                # D, E, F: NA
                na_font = Font(name=FONT_NAME, size=ID_FONT_SIZE, color='999999')
                set_cell(ws, row_num, 4, 'NA', na_font)
                set_cell(ws, row_num, 5, 'NA', na_font)
                set_cell(ws, row_num, 6, 'NA', na_font)

                # 数据行行高: 上游内容按换行/列宽估算
                set_row_height(ws, row_num, [
                    (data.get('up_text', ''), 55),
                ])

                row_num += 1
                seq += 1

    os.makedirs(os.path.dirname(output_path) or '.', exist_ok=True)
    wb.save(output_path)
