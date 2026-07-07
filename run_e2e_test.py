"""
端到端验证脚本 — 不依赖Streamlit，直接运行完整管线并对比基准
"""
import os
import sys
import time
import warnings
warnings.filterwarnings('ignore')

# 确保项目根目录在路径中
sys.path.insert(0, os.path.dirname(os.path.abspath(__file__)))

from config import PROJECT_ROOT, OUTPUT_DIR, MatchCategory
from core.traceability_matrix import build_backward_matrix
from core.text_matcher import verify_matrix
from core.excel_generator import generate_excel


def main():
    # === 1. 配置路径 ===
    downstream_dir = "系统需求"
    upstream_dirs = ["用户需求"]

    downstream_pdf = None
    ds_path = os.path.join(PROJECT_ROOT, downstream_dir)
    for fname in os.listdir(ds_path):
        if fname.lower().endswith('.pdf'):
            downstream_pdf = os.path.join(ds_path, fname)
            break
    if not downstream_pdf:
        print(f"[ERROR] 在 {ds_path} 中未找到PDF"); return

    upstream_pdfs = {}
    for d in upstream_dirs:
        d_path = os.path.join(PROJECT_ROOT, d)
        if not os.path.isdir(d_path):
            continue
        for fname in os.listdir(d_path):
            if fname.lower().endswith('.pdf'):
                doc_name = fname.replace('.pdf', '')
                upstream_pdfs[doc_name] = os.path.join(d_path, fname)

    print(f"下游PDF: {os.path.basename(downstream_pdf)}")
    print(f"上游PDFs: {list(upstream_pdfs.keys())}")
    print()

    # === 2. 构建逆向追踪矩阵 ===
    print(">>> 构建逆向追踪矩阵...")
    t0 = time.time()
    matrix = build_backward_matrix(downstream_pdf, upstream_pdfs)
    print(f"    完成: {len(matrix.rows)} 条关系 ({time.time()-t0:.1f}s)")

    # 检查上游内容填充情况
    empty_up = sum(1 for r in matrix.rows if not r.upstream_content.strip())
    print(f"    上游内容为空: {empty_up}/{len(matrix.rows)}")

    # 打印每行摘要
    print("\n    行号 | 下游ID              | 上游文档            | 上游引用                | 上游内容长度")
    print("    " + "-" * 100)
    for r in matrix.rows:
        ds_short = r.downstream_id[:20]
        doc_short = r.upstream_doc[:18]
        ref_short = r.upstream_ref[:24]
        print(f"    {r.seq_number:>4} | {ds_short:<20} | {doc_short:<18} | {ref_short:<24} | {len(r.upstream_content):>5}")

    # === 3. 执行微块匹配验证 ===
    print("\n>>> 执行微块级文本匹配验证...")
    t1 = time.time()
    verify_matrix(matrix)
    print(f"    完成 ({time.time()-t1:.1f}s)")

    # 统计匹配结果
    green = blue = black = 0
    for r in matrix.rows:
        if r.match_result:
            cat = r.match_result.overall_category
            if cat == MatchCategory.GREEN:
                green += 1
            elif cat == MatchCategory.BLUE:
                blue += 1
            else:
                black += 1
    total = len(matrix.rows)
    print(f"\n    匹配统计:")
    print(f"    GREEN(完全一致): {green}/{total} ({green/total*100:.0f}%)")
    print(f"    BLUE(语义匹配):  {blue}/{total} ({blue/total*100:.0f}%)")
    print(f"    BLACK(未匹配):   {black}/{total} ({black/total*100:.0f}%)")

    # === 4. 生成逆向追踪矩阵Excel ===
    os.makedirs(OUTPUT_DIR, exist_ok=True)
    output_path = os.path.join(OUTPUT_DIR, "追踪验证结果_v5.xlsx")
    print(f"\n>>> 生成逆向追踪矩阵Excel: {output_path}")
    generate_excel(matrix, output_path, include_forward=False)
    print(f"    完成!")

    # === 4b. 生成正向追踪矩阵Excel (独立文件) ===
    forward_output = os.path.join(OUTPUT_DIR, "正向追踪矩阵.xlsx")
    print(f"\n>>> 生成正向追踪矩阵Excel: {forward_output}")
    try:
        from generate_forward_matrix import generate_forward_matrix
        user_pdf_dir = os.path.join(PROJECT_ROOT, "用户需求")
        sys_req_dir = os.path.join(PROJECT_ROOT, downstream_dir)
        fwd_result = generate_forward_matrix(
            output_path, user_pdf_dir, forward_output, sys_req_dir,
        )
        print(f"    文档数: {fwd_result['doc_count']}, "
              f"总条目: {fwd_result['total_entries']}, "
              f"有追踪: {fwd_result['traced']}, 未追踪: {fwd_result['untraced']}, "
              f"系统需求追踪: {fwd_result['total_sys_reqs']}")
    except Exception as e:
        print(f"    [ERROR] 正向矩阵生成失败: {e}")
        import traceback; traceback.print_exc()

    # === 5. 对比基准文件 ===
    baseline_path = os.path.join(PROJECT_ROOT, "基准数据", "Trace_Base.xlsx")
    if os.path.exists(baseline_path):
        print(f"\n>>> 对比基准文件: {baseline_path}")
        _compare_with_baseline(output_path, baseline_path, matrix)
    else:
        print(f"\n[WARN] 基准文件不存在: {baseline_path}")

    print("\n=== 端到端验证完成 ===")


def _compare_with_baseline(result_path, baseline_path, matrix):
    """对比结果与基准文件"""
    try:
        from openpyxl import load_workbook
    except ImportError:
        print("[WARN] openpyxl未安装, 跳过对比"); return

    wb_result = load_workbook(result_path, data_only=True)
    wb_base = load_workbook(baseline_path, data_only=True)

    # 找到逆向追踪矩阵sheet
    sheet_names_r = wb_result.sheetnames
    sheet_names_b = wb_base.sheetnames
    print(f"    结果文件sheets: {sheet_names_r}")
    print(f"    基准文件sheets: {sheet_names_b}")

    # 对比逆向矩阵(第一个sheet)
    ws_r = wb_result[sheet_names_r[0]]
    ws_b = wb_base[sheet_names_b[0]]

    rows_r = ws_r.max_row
    rows_b = ws_b.max_row
    cols_r = ws_r.max_column
    cols_b = ws_b.max_column
    print(f"\n    逆向矩阵对比:")
    print(f"    结果: {rows_r}行 x {cols_r}列")
    print(f"    基准: {rows_b}行 x {cols_b}列")

    # 逐行对比内容(跳过表头)
    mismatches = 0
    for row_idx in range(2, min(rows_r, rows_b) + 1):
        for col_idx in range(1, min(cols_r, cols_b) + 1):
            val_r = str(ws_r.cell(row_idx, col_idx).value or '').strip()
            val_b = str(ws_b.cell(row_idx, col_idx).value or '').strip()
            if val_r != val_b:
                mismatches += 1
                if mismatches <= 10:
                    # 截断显示
                    vr = val_r[:50] + '...' if len(val_r) > 50 else val_r
                    vb = val_b[:50] + '...' if len(val_b) > 50 else val_b
                    print(f"    差异 R{row_idx}C{col_idx}:")
                    print(f"        结果: {vr}")
                    print(f"        基准: {vb}")

    if mismatches > 10:
        print(f"    ... 共 {mismatches} 处差异 (仅显示前10处)")
    elif mismatches == 0:
        print("    完全一致!")
    else:
        print(f"    共 {mismatches} 处差异")

    wb_result.close()
    wb_base.close()


if __name__ == '__main__':
    main()
