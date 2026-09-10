# -*- coding: utf-8 -*-
"""临时脚本: 检查基准数据中非富文本单元格的整体字体颜色/填充 (用后即删)"""
import os
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

BASE_DIR = r"e:\Trace_NL\基准数据"

def cell_font_color(cell):
    f = cell.font
    if f and f.color:
        c = f.color
        if c.type == 'rgb' and c.rgb:
            return str(c.rgb)
        if c.type == 'theme':
            return f"theme:{c.theme}"
        if c.type == 'indexed':
            return f"indexed:{c.indexed}"
    return None

def cell_fill(cell):
    fill = cell.fill
    if fill and fill.fgColor and fill.fgColor.rgb and str(fill.fgColor.rgb) not in ('00000000',):
        return f"fill:{fill.fgColor.rgb}"
    return None

for name in ["Trace_Base.xlsx", "仪控报警.xlsx", "总体方案.xlsx"]:
    p = os.path.join(BASE_DIR, name)
    wb = load_workbook(p, rich_text=True)
    print("=" * 90)
    print(f"FILE: {name}")
    for ws in wb.worksheets:
        for row in ws.iter_rows(min_row=2, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                if isinstance(v, CellRichText):
                    continue  # 富文本已单独分析
                fc = cell_font_color(cell)
                fl = cell_fill(cell)
                s = str(v).replace('\n', '\\n')[:60]
                print(f"  {cell.coordinate} font={fc} {fl or ''} :: {s}")
