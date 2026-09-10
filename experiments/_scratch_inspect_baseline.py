# -*- coding: utf-8 -*-
"""临时脚本: 检查基准数据Excel的富文本标色 (用后即删)"""
import os
from openpyxl import load_workbook
from openpyxl.cell.rich_text import CellRichText

BASE_DIR = r"e:\Trace_NL\基准数据"

def color_of(run):
    f = run.font
    if f is None:
        return None
    c = f.color
    if c is None:
        return None
    if c.type == 'rgb' and c.rgb:
        return str(c.rgb)
    if c.type == 'theme':
        return f"theme:{c.theme}"
    if c.type == 'indexed':
        return f"indexed:{c.indexed}"
    return None

def inspect(path):
    print("=" * 100)
    print(f"FILE: {os.path.basename(path)}")
    wb = load_workbook(path, rich_text=True)
    for ws in wb.worksheets:
        print(f"\n--- SHEET: {ws.title}  dims={ws.dimensions}  max_row={ws.max_row} max_col={ws.max_column} ---")
        if ws.merged_cells.ranges:
            print(f"  merged: {[str(r) for r in ws.merged_cells.ranges]}")
        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=ws.max_column):
            for cell in row:
                v = cell.value
                if v is None:
                    continue
                addr = cell.coordinate
                if isinstance(v, CellRichText):
                    parts = []
                    for seg in v:
                        if isinstance(seg, str):
                            parts.append(("plain", seg))
                        else:
                            parts.append((str(color_of(seg)), seg.text))
                    print(f"  {addr}(rich, {len(parts)} runs):")
                    for col, txt in parts:
                        t = txt.replace('\n', '\\n')
                        print(f"      [{col}] {t!r}")
                else:
                    s = str(v).replace('\n', '\\n')
                    if len(s) > 200:
                        s = s[:200] + '...'
                    print(f"  {addr}: {s}")

for name in ["Trace_Base.xlsx", "仪控报警.xlsx", "总体方案.xlsx"]:
    p = os.path.join(BASE_DIR, name)
    if os.path.exists(p):
        inspect(p)
